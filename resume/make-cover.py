#!/usr/bin/env python3
"""
make-cover.py — Generate a styled cover letter DOCX + PDF from a Markdown file.

Clones cover-template.docx (same directory as this script), replaces paragraph
text while preserving all formatting, then converts to PDF via LibreOffice.

Usage: python3 make-cover.py <cover-letter.md>
"""

import sys
import re
import copy
import subprocess
from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement

SCRIPT_DIR = Path(__file__).parent
TEMPLATE = SCRIPT_DIR / "cover-template.docx"


def parse_cover_letter(path):
    """Parse markdown cover letter into structured sections."""
    content = Path(path).read_text()
    sections = re.split(r'\n\n+', content.strip())

    subject = greeting = None
    body_paras = []

    for section in sections:
        s = section.strip()
        if s.startswith('#') and subject is None:
            subject = re.sub(r'^#+\s*', '', s)
        elif re.match(r'^Dear ', s) and greeting is None:
            greeting = s
        elif s == 'Best regards,':
            break
        elif subject is not None and greeting is not None and s:
            body_paras.append(s)

    if not subject or not greeting or not body_paras:
        raise ValueError(f"Could not parse cover letter structure from {path}")

    return subject, greeting, body_paras


def replace_para_text(para, text):
    """Replace paragraph text, preserving run properties of the first run."""
    if not para.runs:
        para.add_run(text)
        return
    para.runs[0].text = text
    for run in para.runs[1:]:
        run._r.getparent().remove(run._r)


def clone_para_element(source_para):
    """Deep-copy a paragraph XML element (style + formatting, no text)."""
    import copy
    el = copy.deepcopy(source_para._p)
    # Clear run text so we start fresh
    from docx.oxml.ns import qn
    for r in el.findall(qn('w:r')):
        el.remove(r)
    return el


def insert_paragraph_after(doc, ref_para, text, style_para):
    """Insert a new paragraph with style_para's formatting after ref_para."""
    new_p_el = clone_para_element(style_para)
    ref_para._p.addnext(new_p_el)
    # Access the new paragraph via the document body
    from docx.text.paragraph import Paragraph
    new_para = Paragraph(new_p_el, doc)
    new_para.add_run(text)
    # Copy run properties from style_para's first run
    if style_para.runs:
        from docx.oxml.ns import qn
        import copy
        src_rPr = style_para.runs[0]._r.find(qn('w:rPr'))
        if src_rPr is not None:
            new_run_r = new_para.runs[0]._r
            existing_rPr = new_run_r.find(qn('w:rPr'))
            if existing_rPr is not None:
                new_run_r.remove(existing_rPr)
            new_run_r.insert(0, copy.deepcopy(src_rPr))
    return new_para


def make_cover(md_path):
    md_path = Path(md_path)
    docx_path = md_path.with_suffix('.docx')
    pdf_path = md_path.with_suffix('.pdf')

    subject, greeting, body_paras = parse_cover_letter(md_path)

    doc = Document(TEMPLATE)
    paras = doc.paragraphs

    # Template structure:
    # [0] H1 subject  [1] empty  [2] greeting  [3] empty
    # [4..] body paragraphs (each followed by empty)
    # then: empty, "Best regards,", empty, Name, contacts...
    # Find where body ends (first "Best regards," paragraph)
    body_start = 4
    body_end = next(
        (i for i, p in enumerate(paras) if p.text.strip() == 'Best regards,'),
        len(paras)
    )

    # Style template for body paragraphs (use first body para from template)
    body_style_para = paras[body_start]

    # Replace subject and greeting
    replace_para_text(paras[0], subject)
    replace_para_text(paras[2], greeting)

    # Collect existing body+empty paragraph elements to remove
    # Body range is [body_start, body_end)
    paras_to_remove = list(paras[body_start:body_end])

    # Build new body: para, empty, para, empty, ...
    # Insert before the first-to-remove paragraph's position
    insert_after = paras[body_start - 1]  # last empty before body

    new_paras = []
    for i, text in enumerate(body_paras):
        # Insert body paragraph
        p = insert_paragraph_after(doc, insert_after, text, body_style_para)
        new_paras.append(p)
        insert_after = p

        # Insert empty paragraph after (except after last body para)
        if i < len(body_paras) - 1:
            empty_el = clone_para_element(body_style_para)
            p._p.addnext(empty_el)
            from docx.text.paragraph import Paragraph
            empty_p = Paragraph(empty_el, doc)
            empty_p.runs  # just access to ensure it's live
            insert_after = empty_p

    # Remove old body paragraphs
    for p in paras_to_remove:
        p._p.getparent().remove(p._p)

    doc.save(docx_path)
    print(f"  [DOCX]  OK -> {docx_path}")

    result = subprocess.run(
        ['libreoffice', '--headless', '--convert-to', 'pdf',
         str(docx_path), '--outdir', str(md_path.parent)],
        capture_output=True, text=True
    )
    if result.returncode == 0:
        print(f"  [PDF]   OK -> {pdf_path}")
    else:
        print(f"  [PDF]   FAILED\n{result.stderr}")
        sys.exit(1)


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(f"Usage: python3 {sys.argv[0]} <cover-letter.md>")
        sys.exit(1)
    make_cover(sys.argv[1])
